"""Text extraction, overlap audit and external tokenizer comparison. Does NOT train."""
import argparse,collections,hashlib,io,json,re,zipfile
from pathlib import Path
import sentencepiece as spm
from common import ORIGINAL_ZIP,training_pages,forms,alpha_len,paragraphs

BOOKS={'ityala_lamawele','zemkinkomo_magwalandini'}

def docx_text(data):
 from docx import Document
 doc=Document(io.BytesIO(data))
 parts=[p.text for p in doc.paragraphs if p.text.strip()]
 for t in doc.tables:
  for row in t.rows:
   parts.extend(' '.join(c.text for c in row.cells).splitlines())
 return '\n\n'.join(parts)

def iterate_files(name,data,depth=0):
 if depth>2:return
 ext=Path(name).suffix.lower()
 if ext=='.zip':
  with zipfile.ZipFile(io.BytesIO(data)) as z:
   for member in z.infolist():
    if member.is_dir() or member.filename.startswith('__MACOSX/') or member.file_size>100_000_000:continue
    yield from iterate_files(name+'::'+member.filename,z.read(member),depth+1)
 elif ext=='.docx':
  try:yield (name,docx_text(data))
  except Exception as e: print('DOCX read failed',name,str(e)[:80])
 elif ext=='.txt':
  yield (name,data.decode('utf-8',errors='replace'))
 elif ext=='.jsonl':
  for idx,line in enumerate(data.decode('utf-8',errors='replace').splitlines()):
   if not line.strip():continue
   try:r=json.loads(line)
   except json.JSONDecodeError:continue
   if isinstance(r.get('text'),str):
    yield (name+f'::article_{idx:04d}_{r.get("edition","na")}',r['text'])

def documents(root):
 root=Path(root)
 for folder in sorted(root.iterdir()):
  if not folder.is_dir() or folder.name not in BOOKS|{'vukuzenzele_xho'}:continue
  for p in sorted(folder.iterdir()):
   if not p.is_file():continue
   for name,text in iterate_files(p.name,p.read_bytes()):
    if alpha_len(text)>100:
     yield {'id':folder.name+'/'+name,'corpus':folder.name,'text':text}

def audit(trainzip, external, modeldir, outfile):
 d=Path(modeldir);manifest=json.loads((d/'training_manifest.json').read_text())
 precommit=json.loads((d/'selection_frozen.json').read_text())
 train=training_pages(trainzip)
 train_forms=set(w for p in train for w in forms(p['text']))
 alphabet=set(manifest['alphabet'])
 k=12
 # Build exact training paragraph shingles, preserving all blank-line boundaries.
 shingles=set()
 for p in train:
  for block in paragraphs(p['text']):
   ww=forms(block)
   shingles.update(tuple(ww[i:i+k]) for i in range(len(ww)-k+1))
 print('training 12-word paragraph shingles:',len(shingles))
 models={}
 for r in manifest['runs']:
  models[(r['seed'],r['imvo_share'])]=spm.SentencePieceProcessor(model_file=str(d/(r['name']+'.model')))
 bycorp=collections.defaultdict(list)
 docs=list(documents(external))
 if not docs: raise ValueError('No usable external articles/books were downloaded. Run fetch_corpora.py or add the original files.')
 for rec in docs:
  text=rec['text']; words=forms(text)
  total_alpha=sum(map(alpha_len,words)); unseen=[w for w in words if w not in train_forms]
  # Score the exact same words for every tokenizer; exclude any token mapped to <unk>
  # by even one model. Otherwise UNK can artificially reduce apparent token cost.
  unique=set(words)
  eligible={w: (all(not c.isalpha() or c in alphabet for c in w) and
                all(model.unk_id() not in model.encode(w,out_type=int) for model in models.values()))
            for w in unique}
  supported=[w for w in words if eligible[w]]
  supported_unseen=[w for w in unseen if eligible[w]]
  matchwords=0;totwords=0
  for par in paragraphs(text):
   ww=forms(par);covered=set();totwords+=len(ww)
   for i in range(len(ww)-k+1):
    if tuple(ww[i:i+k]) in shingles: covered.update(range(i,i+k))
   matchwords+=len(covered)
  results=[]
  for mix in manifest['mixes']:
   pieces=[]; novpieces=[]
   for seed in manifest['seeds']:
    model=models[(seed,mix)]
    cost=sum(len(model.encode(w,out_type=int)) for w in supported)
    novcost=sum(len(model.encode(w,out_type=int)) for w in supported_unseen)
    pieces.append(100*cost/max(1,sum(map(alpha_len,supported))))
    novpieces.append(100*novcost/max(1,sum(map(alpha_len,supported_unseen))))
   results.append({'imvo_share':mix,'supported_E_by_seed':pieces,'supported_unseen_E_by_seed':novpieces,
                   'supported_E_mean':sum(pieces)/len(pieces),'supported_unseen_E_mean':sum(novpieces)/len(novpieces)})
  bycorp[rec['corpus']].append({'id':rec['id'],'sha256_text':hashlib.sha256(text.encode('utf-8')).hexdigest(),
     'words':len(words),'alpha_chars':total_alpha,'unseen_word_occurrences':len(unseen),
     'unseen_alpha_chars':sum(map(alpha_len,unseen)),
     'common_tokenizer_supported_word_occurrences':len(supported),
     'match12_words':matchwords,'match12_total_words':totwords,'costs':results})
 summary={}
 for corpus,rows in bycorp.items():
  total=sum(r['words'] for r in rows);nover=sum(r['unseen_word_occurrences'] for r in rows)
  pieces={}
  for mix in manifest['mixes']:
   matched=[(r['alpha_chars'], next(c for c in r['costs'] if c['imvo_share']==mix)) for r in rows]
   pieces[str(mix)]={'E_mean_document_weighted':sum(s['supported_E_mean'] for _,s in matched)/len(matched),
                    'unseen_E_mean_document_weighted':sum(s['supported_unseen_E_mean'] for _,s in matched)/len(matched)}
  summary[corpus]={'documents':len(rows),'words':total,'alpha_chars':sum(r['alpha_chars'] for r in rows),
      'novel_surface_occurrence_fraction':nover/total if total else None,
      'matched12_word_fraction':sum(r['match12_words'] for r in rows)/max(1,sum(r['match12_total_words'] for r in rows)),
      'common_tokenizer_supported_word_fraction':sum(r['common_tokenizer_supported_word_occurrences'] for r in rows)/total if total else None,
      'model_costs':pieces}
 output={'status':'external metrics; no downstream LM accuracy claims','selection_precommitted':precommit,
         'caveat':'Books may contain historical newspaper reuse and are only two independent works; modern text has OCR artifacts; only forms encodable by all compared tokenizers (no UNK) scored for E',
         'external_summary':summary,'document_results':dict(bycorp)}
 Path(outfile).parent.mkdir(parents=True,exist_ok=True)
 Path(outfile).write_text(json.dumps(output,indent=2,ensure_ascii=False))
 print(json.dumps(summary,indent=2,ensure_ascii=False)[:6000]);print('Saved',outfile)
if __name__=='__main__':
 p=argparse.ArgumentParser();p.add_argument('--zip',default=ORIGINAL_ZIP);p.add_argument('--external',default='external_raw');p.add_argument('--models',default='models');p.add_argument('--out',default='external_scores.json');a=p.parse_args();audit(a.zip,a.external,a.models,a.out)
